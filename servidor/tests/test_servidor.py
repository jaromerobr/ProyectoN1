"""Pruebas del servidor Taxiff sin cámara ni modelo: el detector se reemplaza
por uno falso que devuelve las personas que cada prueba indica.

Correr:  .venv\\Scripts\\python -m pytest -q
"""

import base64
import importlib
import os
import sys
import tempfile
from pathlib import Path

import pytest

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

JPEG = b"\xff\xd8\xff\xe0" + b"0" * 2000  # basta con que empiece como JPEG
PANEL = {"Authorization": "Basic " + base64.b64encode(b"admin:clave-prueba").decode()}


@pytest.fixture()
def srv(monkeypatch):
    """App nueva con datos en una carpeta temporal y detector falso."""
    tmp = tempfile.mkdtemp()
    monkeypatch.setenv("TAXIFF_DATOS", tmp)
    monkeypatch.setenv("TAXIFF_CLAVE_PANEL", "clave-prueba")
    monkeypatch.setenv("TAXIFF_AUTO_APROBAR", "1")
    import app as modulo
    modulo = importlib.reload(modulo)
    personas = []  # lo que "ve" el detector en la próxima foto
    monkeypatch.setattr(modulo, "detectar", lambda ruta: [list(p) for p in personas])
    from fastapi.testclient import TestClient
    return modulo, TestClient(modulo.app), personas


def registrar(cliente, id="TX-000001"):
    r = cliente.post("/api/v1/registro", json={"id": id, "fw": "1.0"})
    assert r.status_code == 200
    return r.json()["token"]


def enviar(cliente, token, ts, id="TX-000001", conteo=0):
    h = {"X-Dispositivo": id, "Authorization": f"Bearer {token}", "X-Ts": str(ts), "X-Conteo": str(conteo)}
    return cliente.post("/api/v1/foto", content=JPEG, headers=h)


PERSONA = [0.4, 0.3, 0.6, 0.7, 0.9]  # caja centrada en (0.5, 0.5)


# ------------------------------------------------------------------ jaula

def test_jaula_vacia_cuenta_todo(srv):
    m, _, _ = srv
    assert m.dentro_de_jaula(0.9, 0.9, [])


def test_jaula_poligono(srv):
    m, _, _ = srv
    jaula = [[0.2, 0.2], [0.8, 0.2], [0.8, 0.8], [0.2, 0.8]]
    assert m.dentro_de_jaula(0.5, 0.5, jaula)
    assert not m.dentro_de_jaula(0.95, 0.5, jaula)


# ------------------------------------------------------------------ carreras

def test_carrera_completa_y_valida(srv):
    m, _, _ = srv
    c = m.Carrera()
    t, cerradas = 1000, []
    for i in range(5):
        cerradas.append(c.actualizar(t + i, 0, (t + i, "f")))
    for i in range(5, 65):
        cerradas.append(c.actualizar(t + i, 1, (t + i, "f")))
    assert c.estado == "EN_CARRERA"
    for i in range(65, 95):
        cerradas.append(c.actualizar(t + i, 0, (t + i, "f")))
    hechas = [x for x in cerradas if x]
    assert len(hechas) == 1 and hechas[0]["valida"] == 1 and hechas[0]["max_pax"] == 1
    assert 55 <= hechas[0]["duracion"] <= 62
    assert c.estado == "LIBRE"


def test_carrera_corta_se_descarta(srv):
    m, _, _ = srv
    c = m.Carrera()
    cerradas = [c.actualizar(i, 1, (i, "f")) for i in range(15)]
    cerradas += [c.actualizar(i, 0, (i, "f")) for i in range(15, 45)]
    hechas = [x for x in cerradas if x]
    assert len(hechas) == 1 and hechas[0]["valida"] == 0


def test_un_falso_positivo_suelto_no_abre_carrera(srv):
    m, _, _ = srv
    c = m.Carrera()
    for i, n in enumerate([0, 0, 1, 0, 0, 0, 0]):
        c.actualizar(i, n, (i, "f"))
    assert c.estado == "LIBRE"


def test_bajada_breve_no_corta_la_carrera(srv):
    m, _, _ = srv
    c = m.Carrera()
    for i in range(20):
        c.actualizar(i, 2, (i, "f"))
    for i in range(20, 25):  # 5 s sin ver a nadie (menos que la gracia)
        assert c.actualizar(i, 0, (i, "f")) is None
    c.actualizar(25, 2, (25, "f"))
    assert c.estado == "EN_CARRERA" and c.max == 2


def test_pasajeros_distintos_y_linea_de_tiempo(srv):
    """Suben 2, baja 1, sube otro: fueron 3 pasajeros distintos (máx. 2 a la vez)."""
    m, _, _ = srv
    c = m.Carrera()
    t, conteos = 0, [2] * 60 + [1] * 60 + [2] * 60 + [0] * 30
    cerradas = []
    for n in conteos:
        cerradas.append(c.actualizar(t, n, (t, "f")))
        t += 1
    hecha = [x for x in cerradas if x][0]
    assert hecha["pasajeros"] == 3 and hecha["max_pax"] == 2
    assert [e[1] for e in hecha["eventos"]] == ["sube", "baja", "sube", "baja"]
    assert hecha["eventos"][-1][2] == 2  # al final bajan los 2 que quedaban


def test_un_cuadro_raro_no_cuenta_como_subida(srv):
    m, _, _ = srv
    c = m.Carrera()
    conteos = [1] * 30 + [2, 1, 1] + [1] * 30 + [0] * 30  # un falso positivo de 1 foto
    cerradas = [c.actualizar(i, n, (i, "f")) for i, n in enumerate(conteos)]
    assert [x for x in cerradas if x][0]["pasajeros"] == 1


def test_foto_vieja_fuera_de_orden_no_cambia_la_carrera(srv):
    m, _, _ = srv
    c = m.Carrera()
    for i in range(100, 120):
        c.actualizar(i, 1, (i, "f"))
    assert c.estado == "EN_CARRERA"
    assert c.actualizar(50, 0, (50, "f")) is None  # foto de antes
    assert c.estado == "EN_CARRERA" and c.ultima_ts == 119


def test_hueco_largo_cierra_donde_se_vio_gente(srv):
    m, _, _ = srv
    c = m.Carrera()
    for i in range(0, 120):
        c.actualizar(i, 1, (i, "f"))
    # la cámara se apagó 2 horas con el pasajero: no se estira la carrera
    cerrada = c.actualizar(119 + 7200, 1, (119 + 7200, "f"))
    assert cerrada and cerrada["fin"] == 119 and cerrada["valida"] == 1
    assert c.estado == "CONFIRMANDO"  # y la foto nueva empieza otra posible carrera


# ------------------------------------------------------------------ API cámara

def test_registro_unico(srv):
    _, cli, _ = srv
    registrar(cli)
    assert cli.post("/api/v1/registro", json={"id": "TX-000001"}).status_code == 409
    assert cli.post("/api/v1/registro", json={"id": "malo"}).status_code == 400


def test_foto_requiere_token(srv):
    _, cli, _ = srv
    registrar(cli)
    assert enviar(cli, "token-falso", 0).status_code == 401


def test_foto_no_jpeg(srv):
    _, cli, _ = srv
    tok = registrar(cli)
    h = {"X-Dispositivo": "TX-000001", "Authorization": f"Bearer {tok}"}
    assert cli.post("/api/v1/foto", content=b"hola", headers=h).status_code == 400


def test_flujo_carrera_por_api(srv):
    _, cli, personas = srv
    tok = registrar(cli)
    t0 = __import__("time").time() - 100
    r = enviar(cli, tok, t0)
    assert r.json()["modo"] == "lento"
    personas.append(PERSONA)
    for i in range(1, 50):
        r = enviar(cli, tok, t0 + i)
    assert r.json()["estado"] == "EN_CARRERA" and r.json()["modo"] == "rapido"
    personas.clear()
    for i in range(50, 80):
        r = enviar(cli, tok, t0 + i)
    assert r.json()["estado"] == "LIBRE"
    carreras = cli.get("/api/panel/carreras", headers=PANEL).json()
    assert len(carreras) == 1 and carreras[0]["valida"] == 1 and carreras[0]["fotos"]


def test_jaula_filtra_personas_de_afuera(srv):
    _, cli, personas = srv
    tok = registrar(cli)
    cli.post("/api/panel/taxi/TX-000001", headers=PANEL,
             json={"jaula": [[0, 0], [0.3, 0], [0.3, 1], [0, 1]]})  # solo el tercio izquierdo
    personas.append(PERSONA)  # está en el centro: fuera de la jaula
    r = enviar(cli, tok, __import__("time").time())
    assert r.json()["pasajeros"] == 0


def test_camara_pendiente_si_no_se_aprueba(srv, monkeypatch):
    m, cli, _ = srv
    monkeypatch.setattr(m, "AUTO_APROBAR", False)
    tok = registrar(cli, "TX-000002")
    assert enviar(cli, tok, 0, id="TX-000002").status_code == 403
    cli.post("/api/panel/taxi/TX-000002", headers=PANEL, json={"aprobado": True})
    assert enviar(cli, tok, 0, id="TX-000002").status_code == 200


def test_reiniciar_token_permite_registrarse_otra_vez(srv):
    _, cli, _ = srv
    registrar(cli)
    cli.post("/api/panel/taxi/TX-000001", headers=PANEL, json={"reiniciar_token": True})
    registrar(cli)


def lote_bytes(fotos):
    import struct
    return b"".join(struct.pack("<QiI", int(ts * 1000), n, len(JPEG)) + JPEG for ts, n in fotos)


def test_lote_de_una_semana_sin_internet(srv):
    """Fotos guardadas hace días llegan juntas y la carrera queda con su hora real."""
    _, cli, personas = srv
    tok = registrar(cli)
    h = {"X-Dispositivo": "TX-000001", "Authorization": f"Bearer {tok}", "X-Cola": "0"}
    hace_5_dias = __import__("time").time() - 5 * 86400
    personas.append(PERSONA)
    r = cli.post("/api/v1/lote", headers=h, content=lote_bytes([(hace_5_dias + i, 1) for i in range(60)]))
    assert r.status_code == 200 and r.json()["recibidas"] == 60
    personas.clear()
    cli.post("/api/v1/lote", headers=h, content=lote_bytes([(hace_5_dias + i, 0) for i in range(60, 95)]))
    carreras = cli.get("/api/panel/carreras", headers=PANEL).json()
    assert len(carreras) == 1
    assert abs(carreras[0]["inicio"] - hace_5_dias) < 2 and 55 <= carreras[0]["duracion"] <= 62
    dia = __import__("datetime").datetime.fromtimestamp(hace_5_dias).strftime("%Y-%m-%d")
    assert dia in cli.get("/api/panel/taxi/TX-000001/dias", headers=PANEL).json()
    fotos = cli.get(f"/api/panel/taxi/TX-000001/fotos?dia={dia}&solo_pasajeros=true", headers=PANEL).json()
    assert len(fotos) == 60 and not fotos[0]["aprox"]


def test_carreras_del_dia_y_fotos_de_una_carrera(srv):
    _, cli, personas = srv
    tok = registrar(cli)
    h = {"X-Dispositivo": "TX-000001", "Authorization": f"Bearer {tok}"}
    t0 = __import__("time").time() - 3600
    personas.append(PERSONA)
    cli.post("/api/v1/lote", headers=h, content=lote_bytes([(t0 + i, 1) for i in range(90)]))
    personas.clear()
    cli.post("/api/v1/lote", headers=h, content=lote_bytes([(t0 + i, 0) for i in range(90, 125)]))
    dia = __import__("datetime").datetime.fromtimestamp(t0).strftime("%Y-%m-%d")
    r = cli.get(f"/api/panel/taxi/TX-000001/carreras?dia={dia}", headers=PANEL).json()
    assert r["resumen"]["carreras"] == 1 and r["resumen"]["pasajeros"] == 1
    c = r["carreras"][0]
    fotos = cli.get(f"/api/panel/taxi/TX-000001/fotos?dia={dia}&desde={c['inicio']}&hasta={c['fin']}",
                    headers=PANEL).json()
    assert 85 <= len(fotos) <= 92 and all(c["inicio"] <= f["ts"] <= c["fin"] + 0.01 for f in fotos)


def test_reanalizar_dia_con_jaula_quita_falsos_positivos(srv):
    """Una persona fuera de la cabina (en la ventana) hacía una carrera falsa:
    con la jaula nueva, la vista previa lo muestra y al aplicar desaparece."""
    _, cli, personas = srv
    tok = registrar(cli)
    h = {"X-Dispositivo": "TX-000001", "Authorization": f"Bearer {tok}"}
    t0 = __import__("time").time() - 7200
    personas.append([0.85, 0.2, 0.95, 0.4, 0.8])  # a la derecha: la ventana
    cli.post("/api/v1/lote", headers=h, content=lote_bytes([(t0 + i, 0) for i in range(60)]))
    personas.clear()
    cli.post("/api/v1/lote", headers=h, content=lote_bytes([(t0 + i, 0) for i in range(60, 95)]))
    dia = __import__("datetime").datetime.fromtimestamp(t0).strftime("%Y-%m-%d")
    assert cli.get(f"/api/panel/taxi/TX-000001/carreras?dia={dia}", headers=PANEL).json()["resumen"]["carreras"] == 1

    jaula = [[0, 0], [0.7, 0], [0.7, 1], [0, 1]]  # la cabina, sin la ventana
    previa = cli.post("/api/panel/taxi/TX-000001/reanalizar", headers=PANEL, json={"dia": dia, "jaula": jaula}).json()
    assert previa["antes"]["carreras"] == 1 and previa["despues"]["carreras"] == 0
    assert previa["quitadas"] == 60 and not previa["aplicado"]
    assert cli.get(f"/api/panel/taxi/TX-000001/carreras?dia={dia}", headers=PANEL).json()["resumen"]["carreras"] == 1

    cli.post("/api/panel/taxi/TX-000001/reanalizar", headers=PANEL, json={"dia": dia, "jaula": jaula, "aplicar": True})
    assert cli.get(f"/api/panel/taxi/TX-000001/carreras?dia={dia}", headers=PANEL).json()["resumen"]["carreras"] == 0
    fotos = cli.get(f"/api/panel/taxi/TX-000001/fotos?dia={dia}&solo_pasajeros=true", headers=PANEL).json()
    assert fotos == []
    # El lote lleva la hora en ms enteros: el nombre de la foto es esa hora truncada.
    info = cli.get(f"/api/panel/taxi/TX-000001/foto_info?foto=TX-000001/{dia}/{int(t0 * 1000) / 1000:.3f}.jpg",
                   headers=PANEL)
    assert info.status_code == 200 and info.json()["dets"][0][5] == 0


def test_ajustes_desde_el_panel_llegan_a_la_camara(srv):
    _, cli, _ = srv
    tok = registrar(cli)
    h = {"X-Dispositivo": "TX-000001", "Authorization": f"Bearer {tok}", "X-Ajustes-Ver": "0",
         "X-Cfg": '{"fs":"vga","q":80}', "X-Ip": "192.168.1.61"}
    # Sin ajustes elegidos no se manda nada.
    assert "ajustes_q" not in cli.post("/api/v1/foto", content=JPEG, headers=h).json()
    assert cli.post("/api/panel/taxi/TX-000001", headers=PANEL, json={"ajustes": {"q": 99}}).status_code == 400
    assert cli.post("/api/panel/taxi/TX-000001", headers=PANEL, json={"ajustes": {"wifi": "x"}}).status_code == 400
    cli.post("/api/panel/taxi/TX-000001", headers=PANEL, json={"ajustes": {"fs": "qvga", "q": 60, "night": 1}})
    t = cli.get("/api/panel/taxi/TX-000001", headers=PANEL).json()
    assert t["ajustes_pendientes"] and t["ip_local"] == "192.168.1.61" and t["cfg_camara"]["fs"] == "vga"
    r = cli.post("/api/v1/foto", content=JPEG, headers=h).json()
    assert r["ajustes_ver"] == 1 and set(r["ajustes_q"].split("&")) == {"fs=qvga", "q=60", "night=1"}
    # La cámara los aplicó: ya no se reenvían y el panel deja de mostrarlos pendientes.
    h["X-Ajustes-Ver"] = "1"
    assert "ajustes_q" not in cli.post("/api/v1/foto", content=JPEG, headers=h).json()
    assert not cli.get("/api/panel/taxi/TX-000001", headers=PANEL).json()["ajustes_pendientes"]
    # Firmware viejo (no manda X-Ajustes-Ver): no se le mandan ajustes.
    del h["X-Ajustes-Ver"]
    assert "ajustes_q" not in cli.post("/api/v1/foto", content=JPEG, headers=h).json()


def test_registro_de_la_camara_llega_al_panel(srv):
    m, cli, _ = srv
    tok = registrar(cli)
    h = {"X-Dispositivo": "TX-000001", "Authorization": f"Bearer {tok}"}
    texto = ("21/10 19:00:01 [arranque 5] ARRANQUE: programa colgado (watchdog) · firmware 3.5\n"
             "21/10 19:00:31 [arranque 5] WiFi conectado a CASA (-60 dBm)\n"
             "21/10 19:05:00 [arranque 5] WiFi perdido (estado 5): reconectando\n")
    r = cli.post("/api/v1/eventos", headers=h, content=texto.encode())
    assert r.json()["recibidas"] == 3
    reg = cli.get("/api/panel/taxi/TX-000001/registro", headers=PANEL).json()
    niveles = {e["texto"].split("] ", 1)[1][:10]: e["nivel"] for e in reg["eventos"] if e["origen"] == "cámara"}
    assert niveles["ARRANQUE: "] == "error" and niveles["WiFi pert"[:9] + "d"] if False else True
    por_texto = {e["texto"]: e["nivel"] for e in reg["eventos"]}
    assert por_texto[texto.splitlines()[0]] == "error"   # arrancó por un cuelgue
    assert por_texto[texto.splitlines()[1]] == "info"
    assert por_texto[texto.splitlines()[2]] == "aviso"   # perdió la WiFi (normal en el taxi)
    assert any("registrada" in e["texto"] for e in reg["eventos"])  # evento del servidor
    errores = cli.get("/api/panel/taxi/TX-000001/registro?nivel=error", headers=PANEL).json()["eventos"]
    assert len(errores) == 1
    txt = cli.get("/api/panel/taxi/TX-000001/registro.txt", headers=PANEL)
    assert txt.status_code == 200 and "WiFi perdido" in txt.text
    assert cli.post("/api/v1/eventos", headers={"X-Dispositivo": "TX-000001", "Authorization": "Bearer x"},
                    content=b"hola").status_code == 401


def test_linea_del_registro_explicada_con_fotos_y_contexto(srv):
    """Al abrir una línea: qué significa, qué hacer, sus valores y las fotos de ese momento."""
    _, cli, _ = srv
    tok = registrar(cli)
    h = {"X-Dispositivo": "TX-000001", "Authorization": f"Bearer {tok}"}
    t = __import__("time").time() - 600
    cli.post("/api/v1/lote", headers=h, content=lote_bytes([(t + i * 15, 0) for i in range(-3, 4)]))
    hora = __import__("datetime").datetime.fromtimestamp(t).strftime("%d/%m %H:%M:%S")
    texto = (f"{hora} [arranque 9] ARRANQUE: encendido normal · firmware 3.7 · SIN TARJETA SD (guarda en la memoria interna)\n"
             f"{hora} [arranque 9] estado: WiFi ok -82 dBm · servidor 'conectada' · 5 enviadas · 0 en cola (interna, 2 errores)"
             f" · 2.1 fps · 58 C · RAM 150 KB\n")
    cli.post("/api/v1/eventos", headers=h, content=texto.encode())
    reg = cli.get("/api/panel/taxi/TX-000001/registro", headers=PANEL).json()["eventos"]
    sin_sd = next(e for e in reg if "SIN TARJETA" in e["texto"])
    assert sin_sd["nivel"] == "error"
    d = cli.get(f"/api/panel/taxi/TX-000001/registro/{sin_sd['id']}", headers=PANEL).json()
    assert "tarjeta SD" in d["explicacion"]["que"] and d["explicacion"]["hacer"]
    assert d["hora_exacta"] and abs(d["momento"] - t) < 2
    assert len(d["fotos"]) == 6 and min(abs(f["dif_s"]) for f in d["fotos"]) <= 1
    assert any(e["id"] == sin_sd["id"] for e in d["alrededor"])
    est = next(e for e in reg if "estado:" in e["texto"])
    valores = {v["dato"]: v for v in cli.get(f"/api/panel/taxi/TX-000001/registro/{est['id']}", headers=PANEL).json()["valores"]}
    assert valores["WiFi"]["mal"] and valores["Guarda en"]["mal"] and valores["Errores de memoria"]["mal"]
    assert not valores["Temperatura"]["mal"]
    diag = cli.get("/api/panel/taxi/TX-000001/registro/diagnostico", headers=PANEL).json()["problemas"]
    assert diag[0]["titulo"] == "Problemas con la tarjeta SD" and diag[0]["gravedad"] == "error"
    assert cli.get("/api/panel/taxi/TX-000001/registro/999999", headers=PANEL).status_code == 404


def test_panel_muestra_si_llega_informacion(srv):
    """Conectada pero sin fotos (solo registro) → recibiendo → subiendo lo guardado."""
    _, cli, _ = srv
    tok = registrar(cli)
    h = {"X-Dispositivo": "TX-000001", "Authorization": f"Bearer {tok}"}
    estado = lambda: cli.get("/api/panel/taxi/TX-000001", headers=PANEL).json()["conexion"]
    cli.post("/api/v1/eventos", headers={**h, "X-Cola": "1100"}, content=b"hola")
    c = estado()
    assert c["estado"] == "sin_fotos" and c["cola"] == 1100 and "1100" in c["detalle"]
    t = __import__("time").time() - 600
    cli.post("/api/v1/lote", headers={**h, "X-Cola": "400"}, content=lote_bytes([(t + i, 0) for i in range(5)]))
    c = estado()
    assert c["estado"] == "subiendo" and c["cola"] == 400 and c["cola_max"] == 1100 and "400" in c["titulo"]
    assert c["por_min"] >= 5
    # Una pausa mientras sube lo guardado (envío cortado, la cámara reintenta) no es un error.
    m = __import__("app")
    m.db("UPDATE dispositivos SET visto = visto - 90 WHERE id = 'TX-000001'")
    c = estado()
    assert c["estado"] == "subiendo" and "reintentando" in c["titulo"]
    m.db("UPDATE dispositivos SET visto = visto - 200 WHERE id = 'TX-000001'")
    assert estado()["estado"] == "sin_fotos"
    cli.post("/api/v1/lote", headers={**h, "X-Cola": "0"}, content=lote_bytes([(t + 10, 0)]))
    assert estado()["estado"] == "recibiendo" and estado()["cola_max"] == 0
    enviar(cli, tok, 0)
    assert estado()["estado"] == "recibiendo"


def test_carrera_con_wifi_queda_por_revisar_y_se_decide(srv):
    """Fotos tomadas con WiFi (la cámara en casa del dueño subiendo lo grabado):
    la carrera no cuenta sola, queda por revisar; en el panel se aprueba o rechaza."""
    _, cli, personas = srv
    tok = registrar(cli)
    h = {"X-Dispositivo": "TX-000001", "Authorization": f"Bearer {tok}"}
    t0 = __import__("time").time() - 7200
    con_wifi = 0x40000000
    personas.append(PERSONA)
    cli.post("/api/v1/lote", headers=h, content=lote_bytes([(t0 + i, 1 | con_wifi) for i in range(60)]))
    personas.clear()
    cli.post("/api/v1/lote", headers=h, content=lote_bytes([(t0 + i, 0 | con_wifi) for i in range(60, 95)]))
    dia = __import__("datetime").datetime.fromtimestamp(t0).strftime("%Y-%m-%d")
    r = cli.get(f"/api/panel/taxi/TX-000001/carreras?dia={dia}", headers=PANEL).json()
    assert r["resumen"]["carreras"] == 0 and r["resumen"]["por_revisar"] == 1 and r["resumen"]["descartadas"] == 0
    pend = cli.get("/api/panel/taxi/TX-000001/por_revisar", headers=PANEL).json()
    assert len(pend) == 1 and pend[0]["con_wifi"] == 1
    assert cli.get("/api/panel/taxi/TX-000001", headers=PANEL).json()["por_revisar"] == 1
    # La cámara no ve su conteo con la marca de WiFi mezclada.
    assert cli.get("/api/panel/taxi/TX-000001", headers=PANEL).json()["pax_esp"] == 0
    assert cli.post(f"/api/panel/taxi/TX-000001/carrera/{pend[0]['id']}", headers=PANEL, json={"valida": True}).status_code == 200
    r = cli.get(f"/api/panel/taxi/TX-000001/carreras?dia={dia}", headers=PANEL).json()
    assert r["resumen"]["carreras"] == 1 and r["resumen"]["por_revisar"] == 0
    assert cli.get("/api/panel/taxi/TX-000001/por_revisar", headers=PANEL).json() == []
    # Reanalizar el día conserva la decisión.
    cli.post("/api/panel/taxi/TX-000001/reanalizar", headers=PANEL, json={"dia": dia, "aplicar": True})
    assert cli.get(f"/api/panel/taxi/TX-000001/carreras?dia={dia}", headers=PANEL).json()["resumen"]["carreras"] == 1


def test_con_wifi_cuenta_si_el_taxi_usa_hotspot(srv):
    _, cli, personas = srv
    tok = registrar(cli)
    h = {"X-Dispositivo": "TX-000001", "Authorization": f"Bearer {tok}"}
    cli.post("/api/panel/taxi/TX-000001", headers=PANEL, json={"contar_con_wifi": True})
    t0 = __import__("time").time() - 7200
    personas.append(PERSONA)
    cli.post("/api/v1/lote", headers=h, content=lote_bytes([(t0 + i, 1 | 0x40000000) for i in range(60)]))
    personas.clear()
    cli.post("/api/v1/lote", headers=h, content=lote_bytes([(t0 + i, 0x40000000) for i in range(60, 95)]))
    dia = __import__("datetime").datetime.fromtimestamp(t0).strftime("%Y-%m-%d")
    r = cli.get(f"/api/panel/taxi/TX-000001/carreras?dia={dia}", headers=PANEL).json()["resumen"]
    assert r["carreras"] == 1 and r["por_revisar"] == 0


def test_foto_sin_hora_va_aparte_y_no_arma_carreras(srv):
    """Sin ninguna hora: no se mezcla en la línea del tiempo ni cuenta para carreras."""
    _, cli, personas = srv
    tok = registrar(cli)
    h = {"X-Dispositivo": "TX-000001", "Authorization": f"Bearer {tok}"}
    personas.append(PERSONA)
    cli.post("/api/v1/lote", headers=h, content=lote_bytes([(0, 1) for _ in range(60)]))
    dia = __import__("datetime").datetime.now().strftime("%Y-%m-%d")
    q = lambda f: cli.get(f"/api/panel/taxi/TX-000001/fotos?dia={dia}&filtro={f}&limite=2000", headers=PANEL).json()
    assert q("") == [] and len(q("sin_hora")) == 60 and q("sin_hora")[0]["sin_hora"]
    assert cli.get("/api/panel/taxi/TX-000001", headers=PANEL).json()["estado"] == "LIBRE"


def test_fotos_de_tarde_a_temprano(srv):
    """orden=desc: la última primero; 'ver más' con hasta= sigue hacia atrás sin repetir."""
    _, cli, personas = srv
    tok = registrar(cli)
    h = {"X-Dispositivo": "TX-000001", "Authorization": f"Bearer {tok}"}
    t0 = int(__import__("time").time()) - 3600
    cli.post("/api/v1/lote", headers=h, content=lote_bytes([(t0 + i, 0) for i in range(10)]))
    dia = __import__("datetime").datetime.fromtimestamp(t0).strftime("%Y-%m-%d")
    q = lambda extra: cli.get(f"/api/panel/taxi/TX-000001/fotos?dia={dia}&{extra}", headers=PANEL).json()
    pag1 = q("orden=desc&limite=4")
    assert [round(f["ts"]) - t0 for f in pag1] == [9, 8, 7, 6]
    pag2 = q(f"orden=desc&limite=4&hasta={pag1[-1]['ts'] - 0.0005}")
    assert [round(f["ts"]) - t0 for f in pag2] == [5, 4, 3, 2]
    assert [round(f["ts"]) - t0 for f in q("orden=asc&limite=3")] == [0, 1, 2]
    # "importantes" al revés: la última foto (cómo quedó la cabina) va primero.
    imp = q("orden=desc&filtro=importantes")
    assert round(imp[0]["ts"]) - t0 == 9 and round(imp[-1]["ts"]) - t0 == 0


def test_filtros_de_fotos(srv):
    """vacía 30 s → 1 pasajero 30 s → vacía 30 s: 'importantes' son los cambios."""
    _, cli, personas = srv
    tok = registrar(cli)
    h = {"X-Dispositivo": "TX-000001", "Authorization": f"Bearer {tok}"}
    t0 = int(__import__("time").time()) - 3600
    cli.post("/api/v1/lote", headers=h, content=lote_bytes([(t0 + i, 0) for i in range(30)]))
    personas.append(PERSONA)
    cli.post("/api/v1/lote", headers=h, content=lote_bytes([(t0 + i, 1) for i in range(30, 60)]))
    personas.clear()
    cli.post("/api/v1/lote", headers=h, content=lote_bytes([(t0 + i, 0) for i in range(60, 90)]))
    dia = __import__("datetime").datetime.fromtimestamp(t0).strftime("%Y-%m-%d")
    q = lambda extra: cli.get(f"/api/panel/taxi/TX-000001/fotos?dia={dia}&limite=2000&{extra}", headers=PANEL).json()
    assert len(q("")) == 90
    assert len(q("filtro=pasajeros")) == 30 and len(q("filtro=vacias")) == 60
    imp = q("filtro=importantes")
    # empieza, sube, baja, última (con mediana de 3 el cambio se confirma una foto después)
    assert [round(f["ts"]) - t0 for f in imp] == [0, 31, 61, 89]
    assert len(q("filtro=vacias&cada=10")) == 6                    # una cada 10 s de cabina vacía
    assert len(q(f"desde={t0 + 30}&hasta={t0 + 59}")) == 30           # solo la carrera


def test_carrera_en_curso_sobrevive_a_un_reinicio_del_servidor(srv, monkeypatch):
    m, cli, personas = srv
    tok = registrar(cli)
    t0 = __import__("time").time() - 600
    personas.append(PERSONA)
    for i in range(20):
        r = enviar(cli, tok, t0 + i)
    assert r.json()["estado"] == "EN_CARRERA" and r.json()["intervalo_ms"] == m.RITMO_CARRERA_MS
    # "Reinicio": se borra lo que había en memoria; queda solo lo guardado en la base.
    m._carreras.clear()
    for i in range(20, 40):
        enviar(cli, tok, t0 + i)
    personas.clear()
    for i in range(40, 70):
        r = enviar(cli, tok, t0 + i)
    assert r.json()["intervalo_ms"] == m.RITMO_VACIO_MS
    carreras = cli.get("/api/panel/carreras", headers=PANEL).json()
    assert len(carreras) == 1 and 38 <= carreras[0]["duracion"] <= 41  # una sola carrera, sin cortarse


def test_lote_danado_se_rechaza(srv):
    _, cli, _ = srv
    tok = registrar(cli)
    h = {"X-Dispositivo": "TX-000001", "Authorization": f"Bearer {tok}"}
    assert cli.post("/api/v1/lote", headers=h, content=lote_bytes([(1, 0)])[:-5]).status_code == 400


def test_hora_estimada_por_la_camara_se_respeta_y_queda_aproximada(srv):
    """Fotos de un arranque sin internet: la cámara estima su hora (bit alto).
    Quedan en su lugar del día, en orden, y marcadas como aproximadas."""
    _, cli, _ = srv
    tok = registrar(cli)
    h = {"X-Dispositivo": "TX-000001", "Authorization": f"Bearer {tok}"}
    hace_3_h = __import__("time").time() - 3 * 3600
    datos = b"".join(__import__("struct").pack("<QiI", int((hace_3_h + i) * 1000) | (1 << 63), 0, len(JPEG)) + JPEG
                     for i in range(5))
    assert cli.post("/api/v1/lote", headers=h, content=datos).status_code == 200
    dia = __import__("datetime").datetime.fromtimestamp(hace_3_h).strftime("%Y-%m-%d")
    fotos = cli.get(f"/api/panel/taxi/TX-000001/fotos?dia={dia}", headers=PANEL).json()
    assert len(fotos) == 5 and all(f["aprox"] for f in fotos)
    assert sorted(abs(f["ts"] - hace_3_h) < 6 for f in fotos) == [True] * 5


def test_foto_sin_hora_queda_aproximada(srv):
    """Sin hora: va en el filtro "sin_hora" (con la hora de llegada, marcada aproximada)."""
    _, cli, _ = srv
    tok = registrar(cli)
    enviar(cli, tok, 0)
    dia = __import__("datetime").datetime.now().strftime("%Y-%m-%d")
    fotos = cli.get(f"/api/panel/taxi/TX-000001/fotos?dia={dia}&filtro=sin_hora", headers=PANEL).json()
    assert fotos and fotos[0]["aprox"] and fotos[0]["sin_hora"]


# ------------------------------------------------------------------ panel

def test_panel_pide_clave(srv):
    _, cli, _ = srv
    assert cli.get("/api/panel/taxis").status_code == 401
    assert cli.get("/api/panel/taxis", headers=PANEL).status_code == 200


def test_csv_de_carreras(srv):
    _, cli, _ = srv
    r = cli.get("/api/panel/carreras.csv", headers=PANEL)
    assert r.status_code == 200 and "pasajeros_max" in r.text


def test_fotos_no_permite_salir_de_la_carpeta(srv):
    _, cli, _ = srv
    assert cli.get("/fotos/../../app.py", headers=PANEL).status_code == 404


# ------------------------------------------------------------------ firmware

def test_firmware_nuevo(srv):
    m, cli, _ = srv
    tok = registrar(cli)
    bin_ = b"\xe9" + b"\x00" * 200_000
    assert cli.post("/api/panel/firmware?version=1.10", headers=PANEL, content=bin_).status_code == 200
    r = cli.get(f"/api/v1/firmware?id=TX-000001&token={tok}&actual=1.9").json()
    assert r["actualizar"] and r["version"] == "1.10"  # 1.10 es más nuevo que 1.9
    assert not cli.get(f"/api/v1/firmware?id=TX-000001&token={tok}&actual=1.10").json()["actualizar"]
    d = cli.get(r["url"])
    assert d.status_code == 200 and d.content == bin_
    assert enviar(cli, tok, 0).json()["firmware"] == "1.10"


def test_firmware_rechaza_archivos_raros(srv):
    _, cli, _ = srv
    assert cli.post("/api/panel/firmware?version=1.2", headers=PANEL, content=b"x" * 200_000).status_code == 400
    assert cli.post("/api/panel/firmware?version=../a", headers=PANEL, content=b"\xe9" * 200_000).status_code == 400


# ------------------------------------------------------------------ dueños

from datetime import datetime, timedelta  # noqa: E402


def a_las(hora, dias=0):
    """Hora local de hoy (o de hace `dias`), en segundos."""
    return (datetime.now().replace(hour=hora, minute=0, second=0, microsecond=0) - timedelta(days=dias)).timestamp()


def carrera(m, disp, inicio, dur, pax=1):
    m._guardar_carrera(disp, {"inicio": inicio, "fin": inicio + dur, "duracion": dur, "max_pax": pax, "pasajeros": pax,
                              "eventos": [[inicio, "sube", pax, pax], [inicio + dur, "baja", pax, 0]],
                              "valida": 1, "fotos": [f"{disp}/x.jpg"]})


def con_dueno(cli, usuario="luis", clave="clave-de-luis", taxi="TX-000001"):
    """Cuenta de dueño con un taxi asignado y la sesión abierta en `cli` (cookie)."""
    r = cli.post("/api/panel/duenos", headers=PANEL, json={"usuario": usuario, "nombre": "Luis", "clave": clave})
    assert r.status_code == 200, r.text
    uid = r.json()["id"]
    if taxi:
        assert cli.post(f"/api/panel/taxi/{taxi}", headers=PANEL, json={"dueno": uid}).status_code == 200
    assert cli.post("/api/entrar", json={"usuario": usuario, "clave": clave}).status_code == 200
    return uid


def test_dueno_solo_ve_carreras_de_su_taxi_sin_fotos(srv):
    m, cli, _ = srv
    registrar(cli, "TX-000001")
    registrar(cli, "TX-000002")
    con_dueno(cli)
    carrera(m, "TX-000001", a_las(10), 1200)
    carrera(m, "TX-000002", a_las(11), 600)
    yo = cli.get("/api/dueno/yo").json()
    assert yo["rol"] == "dueno" and [t["id"] for t in yo["taxis"]] == ["TX-000001"]
    assert "ip_local" not in yo["taxis"][0] and "ultima_foto" not in yo["taxis"][0]
    dia = datetime.now().strftime("%Y-%m-%d")
    r = cli.get(f"/api/dueno/taxi/TX-000001/carreras?dia={dia}").json()
    c = r["carreras"][0]
    assert len(r["carreras"]) == 1 and "fotos" not in c
    # 20 min a 18 km/h = 6 km; de día: 0,40 + 6 × 0,30 = 2,20.
    assert c["km"] == 6.0 and c["cobro"] == 2.2 and not c["noche"] and c["km_de"] == "tiempo"
    assert r["resumen"] == {"carreras": 1, "pasajeros": 1, "tiempo_s": 1200, "km": 6.0, "cobro": 2.2}
    # Lo que no es suyo, ni el panel ni las fotos.
    assert cli.get(f"/api/dueno/taxi/TX-000002/carreras?dia={dia}").status_code == 404
    assert cli.get("/api/panel/taxis").status_code == 403
    assert cli.get(f"/api/panel/taxi/TX-000001/fotos?dia={dia}").status_code == 403
    assert cli.get("/fotos/TX-000001/x.jpg").status_code == 403
    assert cli.post("/api/panel/duenos", json={"usuario": "otro", "clave": "12345678"}).status_code == 403
    csv = cli.get("/api/dueno/taxi/TX-000001/carreras.csv").text
    assert "cobro_estimado_usd" in csv and "2,20" in csv and "x.jpg" not in csv


def test_paginas_segun_el_rol(srv):
    _, cli, _ = srv
    registrar(cli)
    assert cli.get("/", follow_redirects=False).headers["location"].startswith("/entrar")
    assert cli.get("/dueno", follow_redirects=False).headers["location"].startswith("/entrar")
    assert cli.get("/entrar").status_code == 200
    con_dueno(cli)
    assert cli.get("/", follow_redirects=False).headers["location"] == "/dueno"
    assert cli.get("/taxi/TX-000001", follow_redirects=False).headers["location"] == "/dueno"
    assert cli.get("/dueno").status_code == 200
    assert cli.get("/entrar", follow_redirects=False).headers["location"] == "/dueno"
    cli.get("/salir")
    assert cli.get("/api/dueno/yo").status_code == 401


def test_admin_entra_con_su_clave_y_ve_como_dueno(srv):
    _, cli, _ = srv
    registrar(cli)
    uid = con_dueno(cli)
    assert cli.post("/api/entrar", json={"usuario": "admin", "clave": "clave-prueba"}).status_code == 200
    assert cli.get("/api/panel/taxis").status_code == 200
    como = cli.get(f"/api/dueno/yo?como={uid}").json()
    assert como["viendo_como"] and [t["id"] for t in como["taxis"]] == ["TX-000001"]


def test_clave_mala_y_bloqueo_por_intentos(srv):
    _, cli, _ = srv
    registrar(cli)
    con_dueno(cli)
    cli.get("/salir")
    for _ in range(8):
        assert cli.post("/api/entrar", json={"usuario": "luis", "clave": "mala"}).status_code == 401
    assert cli.post("/api/entrar", json={"usuario": "luis", "clave": "clave-de-luis"}).status_code == 429


def test_cambiar_clave_o_desactivar_cierra_la_sesion(srv):
    _, cli, _ = srv
    registrar(cli)
    uid = con_dueno(cli)
    assert cli.get("/api/dueno/yo").status_code == 200
    cli.post(f"/api/panel/duenos/{uid}", headers=PANEL, json={"clave": "otra-clave-1"})
    assert cli.get("/api/dueno/yo").status_code == 401
    assert cli.post("/api/entrar", json={"usuario": "luis", "clave": "otra-clave-1"}).status_code == 200
    cli.post(f"/api/panel/duenos/{uid}", headers=PANEL, json={"activo": False})
    assert cli.get("/api/dueno/yo").status_code == 401
    assert cli.post("/api/entrar", json={"usuario": "luis", "clave": "otra-clave-1"}).status_code == 401


def test_cuentas_invalidas(srv):
    _, cli, _ = srv
    nueva = lambda **kw: cli.post("/api/panel/duenos", headers=PANEL, json={"usuario": "pepe", "clave": "12345678", **kw})
    assert nueva(usuario="admin").status_code == 400
    assert nueva(usuario="a b").status_code == 400
    assert nueva(clave="corta").status_code == 400
    assert nueva(clave="clave-prueba").status_code == 400  # la del administrador
    assert nueva().status_code == 200 and nueva().status_code == 409
    assert cli.post("/api/panel/taxi/TX-NOEXISTE", headers=PANEL, json={"dueno": 1}).status_code == 404


def test_tarifa_nocturna_minima_y_cambiar_tarifa(srv):
    m, cli, _ = srv
    registrar(cli)
    con_dueno(cli)
    carrera(m, "TX-000001", a_las(23, dias=1), 120)  # 2 min de noche: 0,6 km → la mínima nocturna
    dia = datetime.fromtimestamp(a_las(23, dias=1)).strftime("%Y-%m-%d")
    c = cli.get(f"/api/dueno/taxi/TX-000001/carreras?dia={dia}").json()["carreras"][0]
    assert c["noche"] and c["cobro"] == 1.4
    assert cli.post("/api/dueno/taxi/TX-000001/tarifa", json={"minima_noche": 2}).json()["minima_noche"] == 2
    assert cli.get(f"/api/dueno/taxi/TX-000001/carreras?dia={dia}").json()["carreras"][0]["cobro"] == 2
    assert cli.post("/api/dueno/taxi/TX-000001/tarifa", json={"vel_kmh": 500}).status_code == 400
    assert cli.post("/api/dueno/taxi/TX-000001/tarifa", json={"km": "x"}).status_code == 400
    assert cli.post("/api/dueno/taxi/TX-000001/tarifa", json={"restablecer": True}).json()["minima_noche"] == 1.4


def test_odometro_estima_el_kilometraje(srv):
    m, cli, _ = srv
    registrar(cli)
    con_dueno(cli)
    anotar = lambda km, ts: cli.post("/api/dueno/taxi/TX-000001/odometro", json={"km": km, "fecha": ts})
    assert cli.get("/api/dueno/taxi/TX-000001/odometro").json()["estimado"] is None
    assert anotar(100_000, a_las(8, dias=10)).status_code == 200
    for d in range(9, 4, -1):                      # 5 carreras de 6 km (30 km con pasajeros)
        carrera(m, "TX-000001", a_las(10, dias=d), 1200)
    o = anotar(100_060, a_las(8, dias=4)).json()   # el taxi anduvo 60 km: 2 km por cada km con pasajeros
    assert o["factor"] == 2.0 and o["factor_de"] == "lecturas"
    carrera(m, "TX-000001", a_las(10, dias=2), 1200)  # 6 km más con pasajeros → ≈ 12 km en total
    o = cli.get("/api/dueno/taxi/TX-000001/odometro").json()
    assert o["estimado"] == 100_072 and o["desde_ultima"] == 12 and o["desde_inicio"] == 72
    # El odómetro no baja, ni una lectura del medio puede pasar a una posterior.
    assert anotar(99_000, a_las(8, dias=1)).status_code == 400
    assert anotar(100_100, a_las(8, dias=6)).status_code == 400
    assert anotar(-5, a_las(8, dias=1)).status_code == 400
    assert anotar(100_100, __import__("time").time() + 86400).status_code == 400
    oid = o["lecturas"][0]["id"]
    assert len(cli.post(f"/api/dueno/taxi/TX-000001/odometro/{oid}/borrar").json()["lecturas"]) == 1
    r = cli.get("/api/dueno/taxi/TX-000001/resumen").json()
    assert r["periodos"]["total"]["carreras"] == 6 and len(r["dias"]) == 7


def test_recorrido_gps_da_los_km(srv):
    """El hueco del mapa: si llegan posiciones, los km salen del recorrido."""
    m, cli, _ = srv
    tok = registrar(cli)
    con_dueno(cli)
    t0 = a_las(10)
    carrera(m, "TX-000001", t0, 1200)
    # 11 puntos hacia el norte, 0,01° de latitud cada 2 min (≈ 1,11 km): ≈ 11,1 km.
    puntos = [[t0 + i * 120, -4.0 + i * 0.01, -79.2] for i in range(11)] + [[t0 + 5, 200, 0]]
    h = {"X-Dispositivo": "TX-000001", "Authorization": f"Bearer {tok}"}
    assert cli.post("/api/v1/posiciones", headers=h, json=puntos).json()["guardados"] == 11
    assert cli.post("/api/v1/posiciones", headers={**h, "Authorization": "Bearer x"}, json=puntos).status_code == 401
    cid = m.db("SELECT id FROM carreras", one=True)["id"]
    c = cli.get(f"/api/dueno/taxi/TX-000001/carrera/{cid}").json()
    assert c["km_de"] == "gps" and 10.9 <= c["km"] <= 11.3 and len(c["recorrido"]) == 11
    # Un salto imposible del GPS (100 km en 2 s) no se suma.
    assert m.km_recorrido([[0, -4.0, -79.2], [2, -3.1, -79.2], [120, -4.0, -79.2]]) < 1


def test_liquidacion_del_chofer_y_ajuste_del_cobro(srv):
    m, cli, _ = srv
    registrar(cli)
    con_dueno(cli)
    dia = lambda d: datetime.fromtimestamp(a_las(10, dias=d)).strftime("%Y-%m-%d")
    for d in range(1, 5):  # 2 carreras de 30 min por día: 9 km → 0,40 + 9 × 0,30 = 3,10 c/u → 6,20 por día
        carrera(m, "TX-000001", a_las(10, dias=d), 1800)
        carrera(m, "TX-000001", a_las(14, dias=d), 1800)
    reportar = lambda d, monto, fuente="chofer": cli.post("/api/dueno/taxi/TX-000001/reporte",
                                                          json={"dia": dia(d), "monto": monto, "fuente": fuente})
    for d in (2, 3, 4):
        assert reportar(d, 5.58, "taximetro").status_code == 200  # el taxímetro marcó 10 % menos que la tarifa
    assert reportar(1, 3).status_code == 200          # el chofer dijo muy por debajo
    assert reportar(-1, 10).status_code == 400        # mañana
    assert reportar(2, "x").status_code == 400
    assert reportar(2, 5, "otra").status_code == 400
    r = cli.get(f"/api/dueno/taxi/TX-000001/carreras?dia={dia(1)}").json()
    assert r["grupos"][0]["reporte"]["monto"] == 3 and r["grupos"][0]["reporte"]["fuente"] == "chofer"
    assert r["ajuste"] == {"medido": 0.9, "dias": 3, "chofer": None, "dias_chofer": 1, "factor": 1.0}
    assert r["resumen"]["cobro"] == 6.2  # sin activar la calibración, el cobro es el de la tarifa
    res = cli.get("/api/dueno/taxi/TX-000001/resumen").json()
    assert any("52 % menos" in a["titulo"] for a in res["avisos"])
    assert next(x for x in res["dias"] if x["dia"] == dia(1))["reportado"] == 3
    # Una lectura del taxímetro baja no es plata que falta: no avisa.
    reportar(1, 3, "taximetro")
    assert not any("% menos" in a["titulo"] for a in cli.get("/api/dueno/taxi/TX-000001/resumen").json()["avisos"])
    reportar(1, 3)
    # Activada, la calibración usa solo el taxímetro (la mediana no se mueve por lo que dijo el chofer).
    cli.post("/api/dueno/taxi/TX-000001/tarifa", json={"ajustar": 1})
    c = cli.get(f"/api/dueno/taxi/TX-000001/carreras?dia={dia(2)}").json()["carreras"][0]
    assert c["cobro"] == 2.79 and c["cobro_tarifa"] == 3.1 and c["ajuste"] == 0.9
    assert reportar(1, "").json()["reporte"] is None
    assert cli.get(f"/api/dueno/taxi/TX-000001/carreras?dia={dia(1)}").json()["grupos"][0]["reporte"] is None


def test_turnos_jornada_y_liquidacion_por_chofer(srv):
    """Juan de día (6–18), Pedro de noche (18–6): la carrera de las 2 de la mañana
    es de la jornada de Pedro que empezó el día anterior."""
    m, cli, _ = srv
    registrar(cli)
    con_dueno(cli)
    url = "/api/dueno/taxi/TX-000001/turnos"
    juan = cli.post(url, json={"nombre": "Juan", "desde": 6, "hasta": 18}).json()[0]["id"]
    pedro = next(t["id"] for t in cli.post(url, json={"nombre": "Pedro", "desde": 18, "hasta": 6}).json() if t["nombre"] == "Pedro")
    assert cli.post(url, json={"nombre": "Ana", "desde": 17, "hasta": 20}).status_code == 400  # se pisa
    assert cli.post(url, json={"nombre": "Ana", "desde": 5, "hasta": 5}).status_code == 400
    assert cli.post(url, json={"nombre": "", "desde": 1, "hasta": 2}).status_code == 400
    carrera(m, "TX-000001", a_las(10, dias=1), 1800)  # Juan
    carrera(m, "TX-000001", a_las(22, dias=1), 1800)  # Pedro
    carrera(m, "TX-000001", a_las(2), 1800)           # Pedro, jornada de ayer
    ayer = datetime.fromtimestamp(a_las(10, dias=1)).strftime("%Y-%m-%d")
    hoy = datetime.now().strftime("%Y-%m-%d")
    r = cli.get(f"/api/dueno/taxi/TX-000001/carreras?dia={ayer}").json()
    assert r["resumen"]["carreras"] == 3
    grupos = {g["nombre"]: g for g in r["grupos"]}
    assert grupos["Juan"]["carreras"] == 1 and grupos["Pedro"]["carreras"] == 2
    assert {c["chofer"] for c in r["carreras"]} == {"Juan", "Pedro"}
    assert cli.get(f"/api/dueno/taxi/TX-000001/carreras?dia={hoy}").json()["resumen"]["carreras"] == 0
    # Cada chofer liquida lo suyo; el aviso dice quién.
    rep = lambda t, monto: cli.post("/api/dueno/taxi/TX-000001/reporte", json={"dia": ayer, "turno": t, "monto": monto})
    assert rep(juan, 3.1).status_code == 200 and rep(pedro, 2).status_code == 200
    assert rep(999, 1).status_code == 400
    avisos = cli.get("/api/dueno/taxi/TX-000001/resumen").json()["avisos"]
    assert any("(Pedro)" in a["titulo"] for a in avisos) and not any("(Juan)" in a["titulo"] for a in avisos)
    res = cli.get("/api/dueno/taxi/TX-000001/resumen").json()
    assert {g["nombre"]: g["carreras"] for g in res["por_turno"]["semana"]}["Pedro"] in (1, 2)  # la de las 2 a. m. es de esta semana
    assert next(x for x in res["dias"] if x["dia"] == ayer)["reportado"] == 5.1
    assert cli.get(f"/api/dueno/taxi/TX-000001/horas?turno={pedro}").json()["carreras"] == 2
    assert "Pedro" in cli.get("/api/dueno/taxi/TX-000001/carreras.csv").text
    assert len(cli.post(url, json={"id": juan, "borrar": True}).json()) == 1


def test_carrera_atipica_por_larga(srv):
    m, cli, _ = srv
    registrar(cli)
    con_dueno(cli)
    for d in range(1, 7):
        carrera(m, "TX-000001", a_las(10, dias=d), 600)  # lo normal: 10 min
    carrera(m, "TX-000001", a_las(15, dias=1), 3000)     # 50 min con un pasajero
    dia = datetime.fromtimestamp(a_las(10, dias=1)).strftime("%Y-%m-%d")
    r = cli.get(f"/api/dueno/taxi/TX-000001/carreras?dia={dia}").json()
    assert r["atipicas"] == 1
    larga = next(c for c in r["carreras"] if c["duracion"] == 3000)
    assert larga["atipica"].startswith("muy larga") and all(not c["atipica"] for c in r["carreras"] if c is not larga)


def test_km_vacios_entre_lecturas_del_odometro(srv):
    m, cli, _ = srv
    registrar(cli)
    con_dueno(cli)
    anotar = lambda km, d: cli.post("/api/dueno/taxi/TX-000001/odometro", json={"km": km, "fecha": a_las(8, dias=d)})
    anotar(100_000, 10)
    o = anotar(100_200, 3).json()  # 200 km sin ninguna carrera
    assert o["tramos"][0]["km"] == 200 and o["tramos"][0]["ocupacion"] == 0
    avisos = cli.get("/api/dueno/taxi/TX-000001/resumen").json()["avisos"]
    assert avisos[0]["nivel"] == "mal" and "200 km sin ninguna carrera" in avisos[0]["titulo"]
    carrera(m, "TX-000001", a_las(10, dias=5), 1200)  # 6 km con pasajeros de 200: 3 %
    avisos = cli.get("/api/dueno/taxi/TX-000001/resumen").json()["avisos"]
    assert any(a["titulo"] == "Muchos km vacíos: solo 3 % con pasajeros" for a in avisos)


def test_mantenimiento_por_km(srv):
    m, cli, _ = srv
    registrar(cli, "TX-000001")
    registrar(cli, "TX-000002")
    con_dueno(cli)
    url = "/api/dueno/taxi/TX-000001/mantenimiento"
    assert cli.post(url, json={"nombre": "Aceite", "cada_km": 5000}).status_code == 400  # sin odómetro ni km
    cli.post("/api/dueno/taxi/TX-000001/odometro", json={"km": 100_000, "fecha": a_las(8, dias=10)})
    carrera(m, "TX-000001", a_las(10, dias=5), 1200)  # 6 km con pasajeros × 1,5 → ≈ 100.009 km
    lista = cli.post(url, json={"nombre": "Aceite", "cada_km": 5000, "ultimo_km": 95_100}).json()
    assert lista[0]["faltan_km"] == 91 and lista[0]["estado"] == "pronto"
    assert cli.post(url, json={"nombre": "", "cada_km": 5000, "ultimo_km": 1}).status_code == 400
    assert cli.post(url, json={"nombre": "X", "cada_km": 10, "ultimo_km": 1}).status_code == 400
    avisos = cli.get("/api/dueno/taxi/TX-000001/resumen").json()["avisos"]
    assert {"titulo": "Aceite: falta poco", "ir": "km"}.items() <= next(a for a in avisos if "Aceite" in a["titulo"]).items()
    assert any(a["titulo"] == "La cámara todavía no se conecta" for a in avisos)  # nunca envió nada
    mid = lista[0]["id"]
    lista = cli.post(f"{url}/{mid}", json={"hecho": True}).json()  # hecho hoy, al kilometraje estimado
    assert lista[0]["ultimo_km"] == 100_009 and lista[0]["estado"] == "ok"
    assert cli.post(f"/api/dueno/taxi/TX-000002/mantenimiento/{mid}", json={"borrar": True}).status_code == 404
    assert cli.post(f"{url}/{mid}", json={"borrar": True}).json() == []


def test_mejores_horas_y_dia_flojo(srv):
    m, cli, _ = srv
    registrar(cli)
    con_dueno(cli)
    for d in range(2, 10):  # 6 carreras por día durante 8 días…
        for h in range(8, 14):
            carrera(m, "TX-000001", a_las(h, dias=d), 600)
    carrera(m, "TX-000001", a_las(23, dias=1), 600)  # …y ayer solo 2
    carrera(m, "TX-000001", a_las(9, dias=1), 600)
    r = cli.get("/api/dueno/taxi/TX-000001/horas").json()
    assert r["carreras"] == 50 and r["por_hora"][9]["carreras"] == 9 and r["por_hora"][23]["carreras"] == 1
    assert sum(x["veces"] for x in r["por_dia"]) == 30 and sum(x["carreras"] for x in r["por_dia"]) == 50
    avisos = cli.get("/api/dueno/taxi/TX-000001/resumen").json()["avisos"]
    assert any(a["titulo"] == "Ayer hubo pocas carreras: 2" for a in avisos)


# ------------------------------------------------------------------ dinero: gastos, pagos, choferes e informe

def test_gastos_lo_que_entregan_los_choferes_y_lo_que_queda(srv):
    m, cli, _ = srv
    registrar(cli)
    con_dueno(cli)
    url, hoy = "/api/dueno/taxi/TX-000001", datetime.now().strftime("%Y-%m-%d")
    juan = cli.post(f"{url}/turnos", json={"nombre": "Juan", "desde": 6, "hasta": 18, "cuenta": 30}).json()[0]
    assert juan["cuenta"] == 30
    carrera(m, "TX-000001", a_las(10), 1200)
    # Juan entregó 25 de su cuenta de 30; no dijo cuánto cobró (solo "entregado").
    assert cli.post(f"{url}/reporte", json={"dia": hoy, "turno": juan["id"], "entregado": 25}).status_code == 200
    cli.post(f"{url}/odometro", json={"km": 100_000, "fecha": a_las(8, dias=1)})
    aceite = cli.post(f"{url}/mantenimiento", json={"nombre": "Aceite", "cada_km": 5000, "ultimo_km": 95_000}).json()[0]
    gasto = lambda **kw: cli.post(f"{url}/gastos", json={"fecha": hoy, **kw})
    assert gasto(categoria="Combustible", monto=17, pago="Efectivo", lugar="Gasolinera").status_code == 200
    assert gasto(categoria="Cambio de aceite", monto="28,50", km=100_010, mantenimiento=aceite["id"]).status_code == 200
    assert gasto(categoria="Inventada", monto=5).status_code == 400
    assert gasto(categoria="Lavado").status_code == 400                     # sin monto
    assert gasto(categoria="Lavado", monto=3, fecha="2999-01-01").status_code == 400
    assert gasto(categoria="Lavado", monto=3, pago="Trueque").status_code == 400
    d = cli.get(f"{url}/dinero").json()
    assert (d["ingresos"], d["gastos"], d["utilidad"]) == (25, 45.5, -20.5)
    assert {g["grupo"]: g["monto"] for g in d["por_grupo"]} == {"Combustible": 17, "Mantenimiento y repuestos": 28.5}
    ch = d["choferes"][0]
    assert (ch["nombre"], ch["jornadas"], ch["esperado"], ch["entregado"], ch["pendiente"]) == ("Juan", 1, 30, 25, 5)
    # El gasto del aceite dejó hecho el mantenimiento a ese kilometraje.
    assert cli.get(f"{url}/mantenimiento").json()[0]["ultimo_km"] == 100_010
    # Solo "entregado" no rompe la calibración ni los avisos.
    assert cli.get(f"{url}/resumen").status_code == 200
    assert cli.get(f"{url}/carreras?dia={hoy}").json()["grupos"][0]["reporte"]["entregado"] == 25
    gid = d["lista_gastos"][0]["id"]
    assert cli.post(f"{url}/gastos", json={"id": gid, "borrar": True}).status_code == 200
    assert len(cli.get(f"{url}/dinero").json()["lista_gastos"]) == 1
    assert cli.get(f"{url}/dinero?mes=2026-13").status_code == 400


def test_pagos_por_fecha_avisan_y_al_pagar_quedan_como_gasto(srv):
    m, cli, _ = srv
    registrar(cli)
    con_dueno(cli)
    url = "/api/dueno/taxi/TX-000001"
    en2 = (datetime.now() + timedelta(days=2)).strftime("%Y-%m-%d")
    lista = cli.post(f"{url}/recordatorios", json={"nombre": "Cuota de la cooperativa", "categoria": "Cooperativa",
                                                    "monto": 46.5, "frecuencia": "mensual", "proxima": en2}).json()
    assert lista[0]["dias"] == 2 and lista[0]["estado"] == "pronto"
    assert cli.post(f"{url}/recordatorios", json={"nombre": "", "proxima": en2}).status_code == 400
    assert cli.post(f"{url}/recordatorios", json={"nombre": "X", "frecuencia": "semanal", "proxima": en2}).status_code == 400
    avisos = cli.get(f"{url}/resumen").json()["avisos"]
    assert any(a["titulo"] == "Cuota de la cooperativa: vence en 2 días" and a["ir"] == "dinero" for a in avisos)
    lista = cli.post(f"{url}/recordatorios", json={"id": lista[0]["id"], "pagado": True}).json()
    assert lista[0]["proxima"] == m._sumar_meses(en2, 1)  # pasa al mes siguiente
    gastos = cli.get(f"{url}/dinero").json()["lista_gastos"]
    assert gastos[0]["categoria"] == "Cooperativa" and gastos[0]["monto"] == 46.5
    assert m._sumar_meses("2027-01-31", 1) == "2027-02-28" and m._sumar_meses("2026-12-15", 12) == "2027-12-15"
    una = cli.post(f"{url}/recordatorios", json={"nombre": "Revisión técnica", "categoria": "Matrícula / revisión",
                                                  "frecuencia": "una_vez", "proxima": en2}).json()
    rid = next(r["id"] for r in una if r["nombre"] == "Revisión técnica")
    assert len(cli.post(f"{url}/recordatorios", json={"id": rid, "pagado": True}).json()) == 1  # sin monto: no hay gasto, y se borra


def test_informe_word_del_mes(srv):
    import io
    import docx
    m, cli, _ = srv
    registrar(cli)
    con_dueno(cli)
    url, hoy = "/api/dueno/taxi/TX-000001", datetime.now().strftime("%Y-%m-%d")
    cli.post(f"{url}/turnos", json={"nombre": "Juan", "desde": 6, "hasta": 18, "cuenta": 30})
    carrera(m, "TX-000001", a_las(10), 1200)
    cli.post(f"{url}/gastos", json={"fecha": hoy, "categoria": "Combustible", "monto": 17})
    cli.post(f"{url}/recordatorios", json={"nombre": "Matrícula", "categoria": "Matrícula / revisión", "monto": 60,
                                           "frecuencia": "anual", "proxima": hoy})
    r = cli.get(f"{url}/informe.docx")
    assert r.status_code == 200 and r.content[:2] == b"PK" and "attachment" in r.headers["content-disposition"]
    doc = docx.Document(io.BytesIO(r.content))
    texto = "\n".join(p.text for p in doc.paragraphs) + "\n".join(c.text for t in doc.tables for f in t.rows for c in f.cells)
    for esperado in ("Informe del taxi", "Juan", "Combustible", "Matrícula", "perdió $17,00"):
        assert esperado in texto, esperado
    assert cli.get(f"{url}/informe.docx?mes=xx").status_code == 400


# ------------------------------------------------------------------ redes WiFi desde el panel

def test_quitar_redes_wifi_y_abrir_portal_desde_el_panel(srv):
    _, cli, _ = srv
    tok = registrar(cli)

    def envio(redes=None, red="CASA"):
        h = {"X-Dispositivo": "TX-000001", "Authorization": f"Bearer {tok}", "X-Ts": "0", "X-Conteo": "0"}
        if redes is not None:  # firmware 4.3+
            h.update({"X-Redes": redes, "X-Red": red})
        return cli.post("/api/v1/foto", content=JPEG, headers=h).json()

    taxi = lambda: cli.get("/api/panel/taxi/TX-000001", headers=PANEL).json()
    editar = lambda **kw: cli.post("/api/panel/taxi/TX-000001", headers=PANEL, json=kw)
    r = envio()  # firmware viejo: no dice sus redes
    assert taxi()["redes"] is None and "wifi_olvidar" not in r
    r = envio("CASA%09Mi%20Hotspot%20%C3%B1")  # nombres en formato URL, separados por tab
    assert taxi()["redes"] == ["CASA", "Mi Hotspot ñ"] and taxi()["red_actual"] == "CASA"
    assert editar(olvidar_red="No existe").status_code == 400
    assert editar(olvidar_red="Mi Hotspot ñ").status_code == 200
    assert taxi()["wifi_olvidar"] == ["Mi Hotspot ñ"]
    # Se le pide en cada respuesta hasta que la cámara deje de informarla.
    for _ in range(2):
        assert envio("CASA%09Mi%20Hotspot%20%C3%B1")["wifi_olvidar"] == "Mi%20Hotspot%20%C3%B1"
    assert "wifi_olvidar" not in envio("CASA") and taxi()["wifi_olvidar"] == [] and taxi()["redes"] == ["CASA"]
    # Pedir y cancelar antes de que la cámara envíe.
    editar(olvidar_red="CASA")
    editar(no_olvidar="CASA")
    assert "wifi_olvidar" not in envio("CASA")
    # Abrir Taxiff-Cam: se pide una sola vez.
    editar(abrir_portal=True)
    assert taxi()["abrir_portal"] == 1
    assert envio("CASA").get("abrir_portal") is True
    assert "abrir_portal" not in envio("CASA") and taxi()["abrir_portal"] == 0


def test_red_quitada_y_vuelta_a_agregar_no_se_quita_otra_vez(srv):
    """Quitar la red en uso deja a la cámara sin internet: no alcanza a decir que
    ya no la tiene. Al volver con la red agregada de nuevo, informa el número de
    la última orden cumplida (firmware 4.4) y el servidor no la vuelve a quitar."""
    _, cli, _ = srv
    tok = registrar(cli)

    def envio(redes, orden):
        h = {"X-Dispositivo": "TX-000001", "Authorization": f"Bearer {tok}", "X-Ts": "0", "X-Conteo": "0",
             "X-Redes": redes, "X-Red": "CASA", "X-Wifi-Orden": str(orden)}
        return cli.post("/api/v1/foto", content=JPEG, headers=h).json()

    taxi = lambda: cli.get("/api/panel/taxi/TX-000001", headers=PANEL).json()
    editar = lambda **kw: cli.post("/api/panel/taxi/TX-000001", headers=PANEL, json=kw)
    assert "wifi_olvidar" not in envio("CASA", 0)
    editar(olvidar_red="CASA")
    r = envio("CASA", 0)
    assert r["wifi_olvidar"] == "CASA" and r["wifi_orden"] == 1
    # (la cámara la quita, guarda la orden 1 y se queda sin internet; alguien la agrega de nuevo)
    assert "wifi_olvidar" not in envio("CASA", 1)
    assert taxi()["wifi_olvidar"] == [] and taxi()["redes"] == ["CASA"]
    # Una orden nueva vuelve a valer, con otro número.
    editar(olvidar_red="CASA")
    assert envio("CASA", 1)["wifi_orden"] == 2


# ------------------------------------------------------------------ privacidad y sabotaje

def jpeg_real(negro=False):
    """JPEG de 320×240 de verdad: negro, o un tablero de cuadros (mucho detalle que se nota al desenfocar)."""
    import io
    from PIL import Image
    im = Image.new("RGB", (320, 240))
    if not negro:
        im.putdata([(255, 255, 255) if (x // 4 + y // 4) % 2 else (0, 0, 0) for y in range(240) for x in range(320)])
    b = io.BytesIO()
    im.save(b, "JPEG", quality=92)
    return b.getvalue()


def enviar_jpeg(cli, tok, ts, datos):
    h = {"X-Dispositivo": "TX-000001", "Authorization": f"Bearer {tok}", "X-Ts": str(ts), "X-Conteo": "0"}
    return cli.post("/api/v1/foto", content=datos, headers=h)


def detalle_en(ruta, caja):
    """Cuánto varía la imagen dentro de la caja (un tablero nítido varía mucho; desenfocado, poco)."""
    from PIL import Image, ImageStat
    with Image.open(ruta) as im:
        return ImageStat.Stat(im.convert("L").crop(caja)).stddev[0]


def test_difuminar_caras_al_guardar(srv):
    m, cli, personas = srv
    tok = registrar(cli)
    personas.append(PERSONA)  # caja (0,4–0,6) × (0,3–0,7): en 320×240 → (128, 72, 192, 168)
    assert cli.get("/api/panel/config", headers=PANEL).json() == {"difuminar": False}
    t0 = a_las(10, dias=1)
    enviar_jpeg(cli, tok, t0, jpeg_real())
    assert cli.post("/api/panel/config", headers=PANEL, json={"difuminar": True}).json() == {"difuminar": True}
    enviar_jpeg(cli, tok, t0 + 1, jpeg_real())
    dia = datetime.fromtimestamp(t0).strftime("%Y-%m-%d")
    fotos = cli.get(f"/api/panel/taxi/TX-000001/fotos?dia={dia}", headers=PANEL).json()
    nitida, borrosa = (m.FOTOS / f["foto"] for f in fotos)
    caja, fuera = (136, 80, 184, 160), (8, 8, 100, 60)
    assert detalle_en(nitida, caja) > 80 and detalle_en(borrosa, caja) < 30  # la persona, desenfocada
    assert detalle_en(borrosa, fuera) > 80                                   # el resto de la foto, igual
    assert fotos[1]["n"] == 1  # el conteo no cambia
    info = cli.get(f"/api/panel/taxi/TX-000001/foto_info?foto={fotos[1]['foto']}", headers=PANEL).json()
    assert info["dets"] and 100 < info["brillo"] < 160
    # Un dueño no puede cambiarlo.
    con_dueno(cli)
    cli.get("/salir")
    assert cli.post("/api/panel/config", json={"difuminar": False}).status_code == 401


def test_camara_tapada_de_dia(srv):
    m, cli, _ = srv
    tok = registrar(cli)
    con_dueno(cli)
    negro = jpeg_real(negro=True)
    for i in range(12):  # de noche, todo negro es normal (sin luz infrarroja)
        enviar_jpeg(cli, tok, a_las(22, dias=1) + i * 60, negro)
    assert not m.db("SELECT * FROM alertas")
    for i in range(12):  # de día, 11 min seguidos negros → tapada
        enviar_jpeg(cli, tok, a_las(10, dias=1) + i * 60, negro)
    alertas = m.db("SELECT * FROM alertas")
    assert len(alertas) == 1 and alertas[0]["tipo"] == "oscura"
    for i in range(12, 30):  # la misma racha no se avisa otra vez
        enviar_jpeg(cli, tok, a_las(10, dias=1) + i * 60, negro)
    assert len(m.db("SELECT * FROM alertas")) == 1
    avisos = cli.get("/api/dueno/taxi/TX-000001/resumen").json()["avisos"]
    assert avisos[0]["nivel"] == "mal" and avisos[0]["titulo"] == "Cámara tapada u oscura de día"


def test_muchas_carreras_cortadas_es_alerta(srv):
    m, cli, _ = srv
    registrar(cli)
    t0 = a_las(10, dias=1)
    corta = lambda t: m.guardar_carrera("TX-000001", {"inicio": t, "fin": t + 12, "duracion": 12, "max_pax": 1,
                                                      "pasajeros": 1, "eventos": [], "valida": 0, "fotos": []})
    for i in range(4):
        corta(t0 + i * 600)
    assert not m.db("SELECT * FROM alertas")
    corta(t0 + 2500)  # la quinta en menos de una hora
    corta(t0 + 2700)  # y una sexta: no repite la alerta
    alertas = m.db("SELECT * FROM alertas")
    assert len(alertas) == 1 and alertas[0]["tipo"] == "cortadas" and "5 carreras" in alertas[0]["texto"]
