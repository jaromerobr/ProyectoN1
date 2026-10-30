"""Informe mensual del taxi en Word (.docx) para el dueño: dinero, choferes,
gastos, día por día, mantenimiento y próximos pagos. Lo arma app.py con los
mismos datos que muestra la página del dueño (ver informe_dueno)."""

import io
from datetime import datetime

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

MESES = ("enero", "febrero", "marzo", "abril", "mayo", "junio", "julio", "agosto", "septiembre", "octubre",
         "noviembre", "diciembre")
DIAS = ("lun", "mar", "mié", "jue", "vie", "sáb", "dom")
AZUL = RGBColor(0x0A, 0x4F, 0xB5)
VERDE = RGBColor(0x1F, 0x8A, 0x4C)
ROJO = RGBColor(0xC6, 0x28, 0x28)
GRIS = RGBColor(0x6F, 0x6F, 0x76)


def usd(v):
    """$1.234,56 (como se escribe en Ecuador)."""
    if v is None:
        return "–"
    s = f"{abs(v):,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")
    return ("-$" if v < 0 else "$") + s


def km(v):
    return "–" if v is None else f"{round(v):,} km".replace(",", ".")


def fecha(dia):
    d = datetime.strptime(dia, "%Y-%m-%d")
    return f"{DIAS[d.weekday()]} {d.day:02d}/{d.month:02d}"


def _sombra(celda, color):
    tc = celda._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), color)
    tc.append(shd)


def _texto(celda, texto, negrita=False, color=None, derecha=False, tam=None):
    celda.text = ""
    p = celda.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    if derecha:
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r = p.add_run(str(texto))
    r.bold = negrita
    if color:
        r.font.color.rgb = color
    if tam:
        r.font.size = Pt(tam)


def tabla(doc, encabezado, filas, anchos, numeros=(), total=None):
    """Tabla con encabezado sombreado; `numeros` = índices de columnas alineadas a la derecha."""
    t = doc.add_table(rows=1, cols=len(encabezado))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(encabezado):
        _texto(t.rows[0].cells[i], h, negrita=True, derecha=i in numeros, tam=9.5)
        _sombra(t.rows[0].cells[i], "DCE6F4")
    for fila in filas:
        celdas = t.add_row().cells
        for i, v in enumerate(fila):
            color = None
            if isinstance(v, tuple):  # (texto, color)
                v, color = v
            _texto(celdas[i], v, derecha=i in numeros, color=color, tam=9.5)
    if total:
        celdas = t.add_row().cells
        for i, v in enumerate(total):
            _texto(celdas[i], v, negrita=True, derecha=i in numeros, tam=9.5)
            _sombra(celdas[i], "F1F3F6")
    for fila in t.rows:
        for i, ancho in enumerate(anchos):
            fila.cells[i].width = Cm(ancho)
    doc.add_paragraph()
    return t


def _titulo(doc, texto):
    h = doc.add_heading(texto, level=1)
    for r in h.runs:
        r.font.color.rgb = AZUL
        r.font.size = Pt(14)


def _nota(doc, texto):
    p = doc.add_paragraph()
    r = p.add_run(texto)
    r.font.size = Pt(9)
    r.font.color.rgb = GRIS
    return p


def informe_docx(datos):
    din = datos["dinero"]
    anio, mes = map(int, din["mes"].split("-"))
    doc = Document()
    for s in doc.sections:
        s.top_margin = s.bottom_margin = Cm(1.8)
        s.left_margin = s.right_margin = Cm(2)
    estilo = doc.styles["Normal"]
    estilo.font.name = "Calibri"
    estilo.font.size = Pt(10.5)

    # Portada breve
    t = doc.add_heading(f"Informe del taxi {datos['taxi']}", level=0)
    for r in t.runs:
        r.font.color.rgb = AZUL
    sub = doc.add_paragraph()
    sub.add_run(f"{MESES[mes - 1].capitalize()} de {anio}").bold = True
    if datos["dueno"]:
        sub.add_run(f" · {datos['dueno']}")
    r = sub.add_run(f"\nGenerado por Taxiff el {datetime.now():%d/%m/%Y %H:%M}")
    r.font.size = Pt(9)
    r.font.color.rgb = GRIS

    # 1. Resumen
    _titulo(doc, "1. Resumen del mes")
    queda = din["utilidad"]
    p = doc.add_paragraph()
    r = p.add_run(f"Este mes el taxi {'dejó' if queda >= 0 else 'perdió'} {usd(abs(queda))}")
    r.bold = True
    r.font.size = Pt(13)
    r.font.color.rgb = VERDE if queda >= 0 else ROJO
    p.add_run(f"  (entregaron {usd(din['ingresos'])} y se gastó {usd(din['gastos'])}).")
    c = din["carreras"]
    filas = [
        ("Entregado por los choferes", usd(din["ingresos"])),
        ("Gastos", usd(din["gastos"])),
        ("Queda", (usd(queda), VERDE if queda >= 0 else ROJO)),
        ("Carreras (cámara Taxiff)", f"{c['carreras']} · {c['pasajeros']} pasajeros"),
        ("≈ Cobrado en esas carreras", usd(c["cobro"])),
        ("≈ Km recorridos", f"{km(din['km_total'])} ({km(din['km_con_pax'])} con pasajeros)"),
        ("Gasto por km", usd(din["costo_km"]) if din["costo_km"] is not None else "–"),
        ("Queda por km", usd(din["utilidad_km"]) if din["utilidad_km"] is not None else "–"),
    ]
    tabla(doc, ("Concepto", "Valor"), filas, (8, 7.5), numeros=(1,))

    # 2. Choferes
    if din["choferes"]:
        _titulo(doc, "2. Choferes")
        filas = []
        for ch in din["choferes"]:
            nombre = ch["nombre"] or "Chofer"
            if ch["desde"] is not None:
                nombre += f" ({ch['desde']}:00–{ch['hasta']}:00)"
            falta = ch["pendiente"]
            filas.append((nombre, ch["jornadas"], usd(ch["cuenta"]) if ch["cuenta"] else "–",
                          usd(ch["esperado"]) if ch["esperado"] is not None else "–", usd(ch["entregado"]),
                          (usd(falta), ROJO if falta and falta > 0 else None) if falta is not None else "–",
                          ch["carreras"], usd(ch["cobro"])))
        tabla(doc, ("Chofer", "Jornadas", "Cuenta", "Debía", "Entregó", "Falta", "Carreras", "≈ Cobrado"), filas,
              (4.2, 1.6, 1.6, 1.8, 1.8, 1.8, 1.6, 1.9), numeros=(1, 2, 3, 4, 5, 6, 7))
        _nota(doc, "Debía = cuenta pactada × jornadas trabajadas (días con carreras o con entrega anotada).")

    # 3. Gastos
    _titulo(doc, "3. Gastos")
    if din["por_categoria"]:
        total = din["gastos"] or 1
        filas = [(g["categoria"], g["grupo"], usd(g["monto"]), f"{round(100 * g['monto'] / total)} %")
                 for g in din["por_categoria"]]
        tabla(doc, ("Categoría", "Grupo", "Monto", "%"), filas, (5, 6, 3, 2), numeros=(2, 3),
              total=("Total", "", usd(din["gastos"]), "100 %"))
        filas = [(fecha(g["fecha"]), g["categoria"], g["descripcion"] or "", g["lugar"] or "", g["pago"] or "",
                  km(g["km"]) if g["km"] else "", usd(g["monto"])) for g in reversed(din["lista_gastos"])]
        doc.add_paragraph().add_run("Detalle").bold = True
        tabla(doc, ("Fecha", "Categoría", "Descripción", "Lugar", "Pago", "Odómetro", "Monto"), filas,
              (1.8, 2.8, 4, 2.6, 2.1, 2, 1.8), numeros=(5, 6))
    else:
        _nota(doc, "No se anotaron gastos este mes.")

    # 4. Día por día
    if din["dias"]:
        _titulo(doc, "4. Día por día")
        filas = [(fecha(x["dia"]), x["carreras"], km(x["km"]), usd(x["cobro"]),
                  usd(x["entregado"]) if x["entregado"] is not None else "–") for x in din["dias"]]
        tabla(doc, ("Día", "Carreras", "≈ Km con pasajeros", "≈ Cobrado", "Entregó"), filas, (3, 2.5, 4, 3, 3),
              numeros=(1, 2, 3, 4),
              total=("Total", c["carreras"], km(din["km_con_pax"]), usd(c["cobro"]), usd(din["ingresos"])))

    # 5. Mantenimiento
    _titulo(doc, "5. Mantenimiento")
    odo = datos["odometro"]
    if odo["estimado"] is not None:
        doc.add_paragraph(f"Kilometraje estimado hoy: ≈ {km(odo['estimado'])} "
                          f"(última lectura: {km(odo['ultima']['km'])} el {datetime.fromtimestamp(odo['ultima']['fecha']):%d/%m/%Y}).")
    if datos["mantenimiento"]:
        estados = {"vencido": ("Ya toca", ROJO), "pronto": ("Pronto", RGBColor(0xC2, 0x6A, 0x00)),
                   "ok": ("Al día", VERDE), "sin_km": ("Anotar odómetro", GRIS)}
        filas = []
        for m in datos["mantenimiento"]:
            falta = m["faltan_km"]
            filas.append((m["nombre"], km(m["cada_km"]), km(m["ultimo_km"]), km(m["proximo_km"]),
                          ("se pasó " + km(-falta)) if falta is not None and falta <= 0 else km(falta),
                          estados[m["estado"]]))
        tabla(doc, ("Tarea", "Cada", "Última vez", "Próxima", "Faltan", "Estado"), filas, (4.2, 2.3, 2.6, 2.6, 2.6, 2.6),
              numeros=(1, 2, 3, 4))
    else:
        _nota(doc, "Sin tareas de mantenimiento anotadas (se agregan en la pestaña Taxi de la página).")

    # 6. Próximos pagos
    _titulo(doc, "6. Próximos pagos y vencimientos")
    pagos = [p for p in datos["recordatorios"] if p["dias"] <= 60]
    if pagos:
        frec = {"mensual": "Cada mes", "anual": "Cada año", "una_vez": "Una vez"}
        filas = []
        for p in pagos:
            n = p["dias"]
            cuando = (f"venció hace {-n} día{'s' if n < -1 else ''}" if n < 0 else "hoy" if n == 0
                      else "mañana" if n == 1 else f"en {n} días")
            filas.append((p["nombre"], f"{fecha(p['proxima'])}/{p['proxima'][:4]}", (cuando, ROJO if n < 0 else None),
                          usd(p["monto"]) if p["monto"] else "–", frec.get(p["frecuencia"], "")))
        tabla(doc, ("Qué", "Vence", "Cuándo", "Monto", "Se repite"), filas, (5.5, 2.8, 3.2, 2.5, 2.5), numeros=(3,))
    else:
        _nota(doc, "Nada vence en los próximos 60 días.")

    _nota(doc, "≈ = estimado. Las carreras las cuenta la cámara Taxiff. Sin GPS, los km de cada carrera salen de su "
               f"duración × {datos['tarifa']['vel_kmh']} km/h y el cobro, de la tarifa anotada; los km totales, de los km con "
               f"pasajeros × {str(din['factor_km']).replace('.', ',')} (según las lecturas del odómetro). "
               "Lo entregado y los gastos son los que anotó el dueño.")
    salida = io.BytesIO()
    doc.save(salida)
    return salida.getvalue()
